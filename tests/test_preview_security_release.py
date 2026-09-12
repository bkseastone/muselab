"""Untrusted preview inputs must keep parser and credential boundaries."""
import csv
import io
import zipfile

import openpyxl
import pytest
from fastapi import HTTPException


def test_legacy_raw_redirect_removes_global_token(client, temp_root):
    from tests.conftest import TEST_TOKEN

    (temp_root / "fixture.svg").write_text('<svg xmlns="http://www.w3.org/2000/svg"/>')
    response = client.get(
        "/api/files/raw", params={"path": "fixture.svg", "token": TEST_TOKEN},
        follow_redirects=False,
    )
    assert response.status_code == 303
    assert TEST_TOKEN not in response.headers["location"]
    assert "ticket=preview." in response.headers["location"]
    assert response.headers["referrer-policy"] == "no-referrer"
    assert response.headers["cache-control"] == "no-store"
    resource = client.get(response.headers["location"])
    assert resource.status_code == 200
    assert "sandbox allow-scripts" in resource.headers["content-security-policy"]


def test_preview_ticket_remains_reusable_and_file_bound(client, auth, temp_root):
    (temp_root / "fixture.pdf").write_bytes(b"%PDF-synthetic-resource")
    ticket = client.post(
        "/api/files/preview-ticket", headers=auth, json={"path": "fixture.pdf"},
    ).json()["ticket"]
    for _ in range(2):
        result = client.get("/api/files/raw", params={"path": "fixture.pdf", "ticket": ticket})
        assert result.status_code == 200
    assert client.get("/api/files/raw", params={"path": "README.md", "ticket": ticket}).status_code == 401


def test_xlsx_preview_rejects_compression_budget_before_parser(client, auth, temp_root, monkeypatch):
    path = temp_root / "compressed.xlsx"
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("xl/sharedStrings.xml", b"x" * (2 * 1024 * 1024))

    def should_not_parse(*args, **kwargs):
        pytest.fail("unsafe archive reached openpyxl")

    monkeypatch.setattr(openpyxl, "load_workbook", should_not_parse)
    response = client.get("/api/files/xlsx?path=compressed.xlsx", headers=auth)
    assert response.status_code == 422
    assert "compression ratio" in response.json()["detail"]


def test_shared_xlsx_budget_accepts_bytes_and_paths(tmp_path):
    from backend.spreadsheet_safety import validate_xlsx_archive

    source = io.BytesIO()
    with zipfile.ZipFile(source, "w") as archive:
        archive.writestr("sheet.xml", b"ordinary fixture")
    path = tmp_path / "sample.xlsx"
    path.write_bytes(source.getvalue())
    for value in (source.getvalue(), path):
        validate_xlsx_archive(value)
        with pytest.raises(HTTPException, match="entry budget"):
            validate_xlsx_archive(value, max_entries=0)


def test_xlsx_preview_limits_iterator_without_padding_short_sheets(client, auth, temp_root, monkeypatch):
    from openpyxl.worksheet._read_only import ReadOnlyWorksheet

    workbook = openpyxl.Workbook()
    workbook.active["A1"] = "ordinary cell"
    workbook.save(temp_root / "small.xlsx")
    workbook.close()
    original = ReadOnlyWorksheet.iter_rows
    observations = []

    def observe(self, *args, **kwargs):
        observations.append(kwargs)
        return original(self, *args, **kwargs)

    monkeypatch.setattr(ReadOnlyWorksheet, "iter_rows", observe)
    result = client.get("/api/files/xlsx?path=small.xlsx", headers=auth)
    assert result.status_code == 200
    assert result.json()["sheets"][0]["rows"] == [["ordinary cell"]]
    assert observations == [{"max_row": 1, "max_col": 1, "values_only": True}]


def test_csv_large_field_returns_actionable_422(client, auth, temp_root):
    (temp_root / "large-field.csv").write_text("name,data\nfixture," + "x" * (csv.field_size_limit() + 1))
    result = client.get("/api/files/csv?path=large-field.csv", headers=auth)
    assert result.status_code == 422
    assert "field" in result.json()["detail"]
    assert "fixture" not in result.json()["detail"]
